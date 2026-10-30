"""Create fictional Excel workbooks for the TPA dashboard portfolio demo.

Place this script in src/ and run: python src/create_sample_data.py
In sync_balances.py, use .xlsx filenames for the three demo workbooks.
Existing files are never overwritten. No database or history is created here.
"""
from pathlib import Path
from openpyxl import Workbook

BASE = Path(__file__).resolve().parent
ROOT = BASE / "sample_data"
EXAMPLES = {
    "Insurance Company A": 50000.00,
    "Insurance Company B": 75000.00,
    "Insurance Company C": 100000.00,
}


def main():
    targets = {
        company: ROOT / f"Intranet - {company}" / f"{company}.xlsx"
        for company in EXAMPLES
    }
    existing = [path for path in targets.values() if path.exists()]
    if existing:
        raise SystemExit(
            "No files created: sample workbook(s) already exist. "
            "Use a fresh demo folder to regenerate the examples."
        )

    for company, balance in EXAMPLES.items():
        path = targets[company]
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        try:
            ws = wb.active
            ws.title = company
            ws["A1"] = "FICTIONAL SAMPLE DATA — PORTFOLIO DEMO ONLY"
            ws["A2"] = "Invented balances; no real company or financial records."
            ws["D4"] = "Remaining balance"
            ws["E4"] = balance
            ws["E4"].number_format = "#,##0.00"
            ws.column_dimensions["A"].width = 60
            ws.column_dimensions["D"].width = 22
            ws.column_dimensions["E"].width = 18
            # Exclusive creation protects against an existing file appearing
            # after the initial check. These are ordinary .xlsx workbooks.
            with path.open("xb") as output:
                wb.save(output)
        finally:
            wb.close()
        print(f"Created fictional workbook: {path.relative_to(BASE)}")

    print("Done. Run sync_balances.py next to load the fictional balances.")


if __name__ == "__main__":
    main()
